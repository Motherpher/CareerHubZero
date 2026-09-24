from __future__ import annotations

import json
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from dateutil import parser as dtparser
from docx import Document
from docx.shared import Cm, Pt

try:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
except ImportError:  # pragma: no cover
    A4 = None


LEDGER_SCHEMA_VERSION = "1.0"


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def safe_name(value: str) -> str:
    value = re.sub(r"[^\w\- ]+", "", str(value or "")).strip()
    return re.sub(r"\s+", "_", value)[:80] or "hrdm"


def read_ledger(path: Path) -> dict:
    if not path.exists():
        return {"schema_version": LEDGER_SCHEMA_VERSION, "updated_at": None, "runs": []}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        data = {}
    data.setdefault("schema_version", LEDGER_SCHEMA_VERSION)
    data.setdefault("runs", [])
    return data


def write_ledger(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    data["updated_at"] = utcnow()
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def deadline_passed(value: str | None, now: datetime | None = None) -> bool:
    if not value:
        return False
    now = now or datetime.now(timezone.utc)
    try:
        parsed = dtparser.parse(str(value))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.date() < now.date()
    except Exception:
        return False


def _list_lines(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(x) if not isinstance(x, dict) else json.dumps(x, ensure_ascii=False) for x in value]
    if isinstance(value, dict):
        return [f"{k}: {json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v}" for k, v in value.items()]
    return [str(value)]


def report_markdown(hrdm: dict) -> str:
    job = hrdm.get("job") or {}
    hy = hrdm.get("hybridianesque") or {}
    trace = hrdm.get("trace") or {}
    lines = [
        f"# HRDM-R - {job.get('title') or 'Role'}",
        "",
        f"**Arbetsgivare:** {job.get('company') or '-'}  ",
        f"**Plats:** {job.get('location') or '-'}  ",
        f"**Sista ansökningsdag:** {job.get('deadline') or '-'}  ",
        f"**Process-ID:** {hrdm.get('process_id') or trace.get('process_id') or '-'}",
        "",
        "## 1. Signaler",
        "",
    ]
    for item in hrdm.get("signals") or []:
        lines.append(f"- **{item.get('signal','')}** - {item.get('evidence','')} ({item.get('confidence','')})")
    lines += ["", "## 2. Nyckelord och begrepp", ""]
    for cluster in hrdm.get("concept_clusters") or []:
        lines.append(f"- **{cluster.get('cluster','')}**: {', '.join(cluster.get('items') or [])}")
    lines += ["", "## 3. Dolt behov", ""]
    hidden = hrdm.get("hidden_need") or {}
    lines.append(hidden.get("statement") or "-")
    for row in hidden.get("rationale") or []:
        lines.append(f"- {row}")
    lines += ["", "## 4. Fältlogik", ""]
    for row in _list_lines(hrdm.get("field_logic")):
        lines.append(f"- {row}")
    lines += ["", "## 5. FunctionCore", ""]
    fc = hrdm.get("function_core") or {}
    lines.append(f"**Kedja:** {fc.get('chain') or '-'}")
    lines.append("")
    lines.append(fc.get("summary") or "-")
    lines += ["", "## 6. DoD", ""]
    dod = hrdm.get("dod") or {}
    lines.append(f"- Alpha: {dod.get('alpha','')}")
    lines.append(f"- Zenith: {dod.get('zenith','')}")
    lines.append(f"- Klassificering: {dod.get('classification','')}")
    lines.append(f"- Medel: {dod.get('average','')}")
    for k, v in (dod.get("dimensions") or {}).items():
        lines.append(f"- {k}: {v}")
    lines += ["", "## 7. Sannolika bedömningszoner", ""]
    for zone in hrdm.get("assessment_zones") or []:
        lines.append(f"- **{zone.get('zone','')}** ({zone.get('importance','')}): {', '.join(zone.get('evidence_from_ad') or [])}")
    lines += ["", "## 8. Kandidatpositionering", ""]
    cp = hrdm.get("candidate_positioning") or {}
    for key, title in [
        ("strong_matches", "Starka träffar"),
        ("transferable_matches", "Överförbara träffar"),
        ("gaps_unknowns", "Gap / okänt"),
        ("prohibited_claims", "Får inte överdrivas"),
        ("proof_points", "Belägg"),
        ("cv_emphasis", "CV-fokus"),
    ]:
        lines.append(f"### {title}")
        for row in cp.get(key) or []:
            lines.append(f"- {row}")
        lines.append("")
    lines += ["## Hybridianesque (Hy-Filter)", ""]
    lines.append(f"**Rekommenderad:** {'Ja' if hy.get('recommended') else 'Nej'}  ")
    lines.append(f"**Status:** {hy.get('status') or '-'}  ")
    lines.append(f"**Användarsvar:** {hy.get('user_decision') or '-'}")
    lines.append("")
    if hy.get("why"):
        lines.append(hy.get("why"))
        lines.append("")
    for key, label in [
        ("participating_logics", "Deltagande logiker"),
        ("asymmetries", "Asymmetrier"),
        ("translation_relations", "Översättning / medling"),
        ("process_to_output", "Process till output"),
        ("overload_or_misframing", "Överbelastning / felinramning"),
        ("failure_modes", "Failure modes"),
        ("risk_remaining", "Kvarstående risk"),
    ]:
        if hy.get(key):
            lines.append(f"### {label}")
            for row in hy.get(key) or []:
                lines.append(f"- {row}")
            lines.append("")
    if hy.get("framing_if_active"):
        lines += ["### Inramning om filtret är aktivt", "", hy["framing_if_active"], ""]
    lines += ["## 9. HCC-kommentar", ""]
    hcc = hrdm.get("hcc") or {}
    for risk in hcc.get("risks") or []:
        lines.append(f"- {risk}")
    lines.append("")
    lines.append(hcc.get("commentary") or "-")
    lines += ["", "## Ansökningsstrategi", ""]
    strategy = hrdm.get("application_strategy") or {}
    for row in _list_lines(strategy):
        lines.append(f"- {row}")
    lines += ["", "## Trace / ledger", ""]
    for key in [
        "run_id", "timestamp", "environment", "piusite_usage", "hy_filter_usage",
        "hcc_authority", "final_output_state", "process_id_validation_status",
        "last_pid_rev", "runseq"
    ]:
        lines.append(f"- {key}: {trace.get(key)}")
    if trace.get("outstanding_warnings"):
        lines += ["", "### Outstanding warnings"]
        for row in trace["outstanding_warnings"]:
            lines.append(f"- {row}")
    if trace.get("bank_eligible_outputs"):
        lines += ["", "### Bank-eligible outputs"]
        for row in trace["bank_eligible_outputs"]:
            lines.append(f"- {row}")
    return "\n".join(lines).strip() + "\n"


def _docx_add_value(doc: Document, value: Any) -> None:
    if isinstance(value, dict):
        for k, v in value.items():
            p = doc.add_paragraph()
            p.add_run(str(k).replace("_", " ").title() + ": ").bold = True
            p.add_run(json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else str(v))
    elif isinstance(value, list):
        for item in value:
            doc.add_paragraph(json.dumps(item, ensure_ascii=False) if isinstance(item, dict) else str(item), style="List Bullet")
    else:
        doc.add_paragraph(str(value or ""))


def write_hrdm_docx(path: Path, hrdm: dict) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    job = hrdm.get("job") or {}
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(f"HRDM-R - {job.get('title') or 'Role'}", level=0)
    doc.add_paragraph(f"{job.get('company') or ''} | {job.get('location') or ''}")
    doc.add_paragraph(f"Process-ID: {hrdm.get('process_id') or ''}")
    sections = [
        ("Signals", hrdm.get("signals")),
        ("Key Words & Concepts", hrdm.get("concept_clusters")),
        ("Hidden Need Reconstruction", hrdm.get("hidden_need")),
        ("Field Logic Reconstruction", hrdm.get("field_logic")),
        ("FunctionCore Estimation", hrdm.get("function_core")),
        ("DoD Diagnostic", hrdm.get("dod")),
        ("Likely Assessment Zones", hrdm.get("assessment_zones")),
        ("Candidate Positioning Map", hrdm.get("candidate_positioning")),
        ("Hybridianesque (Hy-Filter)", hrdm.get("hybridianesque")),
        ("HCC Reverse Commentary", hrdm.get("hcc")),
        ("Application Strategy", hrdm.get("application_strategy")),
        ("Trace / Run Registry", hrdm.get("trace")),
    ]
    for heading, value in sections:
        doc.add_heading(heading, level=1)
        _docx_add_value(doc, value)
    doc.save(path)
    return path


def write_hrdm_pdf(path: Path, hrdm: dict) -> Path:
    if A4 is None:
        raise RuntimeError("reportlab is required for HRDM PDF export")
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("CareerHubBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, alignment=TA_LEFT, spaceAfter=5)
    h1 = ParagraphStyle("CareerHubH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, spaceAfter=8)
    h2 = ParagraphStyle("CareerHubH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=7, spaceAfter=4)
    story = []
    for raw in report_markdown(hrdm).splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 0.12 * cm))
            continue
        safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("**", "")
        if line.startswith("# "):
            story.append(Paragraph(safe[2:], h1))
        elif line.startswith("## "):
            story.append(Paragraph(safe[3:], h2))
        elif line.startswith("### "):
            story.append(Paragraph(safe[4:], h2))
        elif line.startswith("- "):
            story.append(Paragraph("• " + safe[2:], body))
        else:
            story.append(Paragraph(safe, body))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=1.8 * cm, leftMargin=1.8 * cm, topMargin=1.7 * cm, bottomMargin=1.7 * cm, title=f"HRDM-R {((hrdm.get('job') or {}).get('title') or '')}")
    doc.build(story)
    return path


def save_report_bundle(root: Path, hrdm: dict, ledger_path: Path) -> dict:
    job = hrdm.get("job") or {}
    job_id = str(job.get("id") or safe_name(job.get("url") or job.get("title") or "job"))
    process_id = str(hrdm.get("process_id") or (hrdm.get("trace") or {}).get("process_id") or "")
    if not process_id:
        raise ValueError("HRDM result requires process_id before report banking")
    base = root / "reports" / "hrdm" / "active" / safe_name(job_id) / safe_name(process_id)
    base.mkdir(parents=True, exist_ok=True)

    result_json = base / "HRDM_result.json"
    md_path = base / "HRDM_report.md"
    docx_path = base / "HRDM_report.docx"
    pdf_path = base / "HRDM_report.pdf"

    result_json.write_text(json.dumps(hrdm, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    md_path.write_text(report_markdown(hrdm), encoding="utf-8")
    write_hrdm_docx(docx_path, hrdm)
    write_hrdm_pdf(pdf_path, hrdm)

    ledger = read_ledger(ledger_path)
    run_id = (hrdm.get("trace") or {}).get("run_id")
    entry = {
        "run_id": run_id,
        "process_id": process_id,
        "job_id": job_id,
        "title": job.get("title") or "",
        "company": job.get("company") or "",
        "location": job.get("location") or "",
        "deadline": job.get("deadline") or "",
        "job_url": job.get("url") or "",
        "created_at": (hrdm.get("trace") or {}).get("timestamp") or utcnow(),
        "status": "active",
        "archived_at": None,
        "archive_reason": None,
        "hy_filter": {
            "recommended": bool((hrdm.get("hybridianesque") or {}).get("recommended")),
            "status": (hrdm.get("hybridianesque") or {}).get("status"),
            "used": bool((hrdm.get("trace") or {}).get("hy_filter_usage")),
        },
        "trace": hrdm.get("trace") or {},
        "paths": {
            "markdown": str(md_path.relative_to(root)),
            "docx": str(docx_path.relative_to(root)),
            "pdf": str(pdf_path.relative_to(root)),
            "json": str(result_json.relative_to(root)),
        },
    }
    runs = ledger.setdefault("runs", [])
    old = next((x for x in runs if x.get("run_id") == run_id and run_id), None)
    if old:
        old.update(entry)
    else:
        runs.append(entry)
    write_ledger(ledger_path, ledger)
    return entry


def archive_expired_reports(root: Path, ledger_path: Path, now: datetime | None = None) -> int:
    now = now or datetime.now(timezone.utc)
    ledger = read_ledger(ledger_path)
    changed = 0
    for entry in ledger.get("runs", []):
        if entry.get("status") != "active" or not deadline_passed(entry.get("deadline"), now):
            continue
        old_paths = entry.get("paths") or {}
        process_id = safe_name(entry.get("process_id") or entry.get("run_id") or "run")
        dest = root / "reports" / "hrdm" / "history" / str(now.year) / process_id
        dest.mkdir(parents=True, exist_ok=True)

        new_paths = {}
        for key, rel in old_paths.items():
            source = root / rel
            if source.exists():
                target = dest / source.name
                if target.resolve() != source.resolve():
                    shutil.move(str(source), str(target))
                new_paths[key] = str(target.relative_to(root))
            else:
                new_paths[key] = rel

        entry["status"] = "archived"
        entry["archived_at"] = now.isoformat()
        entry["archive_reason"] = "deadline_passed"
        entry["paths"] = new_paths
        changed += 1

    if changed:
        write_ledger(ledger_path, ledger)
    return changed


def render_report_indexes(root: Path, ledger_path: Path) -> tuple[Path, Path]:
    ledger = read_ledger(ledger_path)
    active = [x for x in ledger.get("runs", []) if x.get("status") == "active"]
    active.sort(key=lambda x: (x.get("deadline") or "9999", x.get("created_at") or ""))

    active_path = root / "HRDM_REPORTS.md"
    hist_path = root / "HRDM_LEDGER.md"

    lines = [
        "# HRDM-analyser",
        "",
        f"**{len(active)} aktiva rapporter**",
        "",
        "Rapporter visas här så länge jobbets sista ansökningsdag inte har passerat. När datumet passerar flyttas rapporten automatiskt till HRDM-ledgern.",
        "",
        "| Roll | Arbetsgivare | Sista dag | Hy-Filter | Läs | Word | PDF |",
        "|---|---|---|---|---|---|---|",
    ]
    for x in active:
        p = x.get("paths") or {}
        hy = (x.get("hy_filter") or {}).get("status") or "-"
        lines.append(f"| {x.get('title','')} | {x.get('company','')} | {str(x.get('deadline') or '-')[:10]} | {hy} | [Rapport]({p.get('markdown','#')}) | [Word]({p.get('docx','#')}) | [PDF]({p.get('pdf','#')}) |")
    if not active:
        lines.append("| - | - | - | - | Inga aktiva rapporter | - | - |")
    active_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    hlines = [
        "# HRDM-ledger",
        "",
        f"**{len(ledger.get('runs', []))} registrerade HRDM-körningar**",
        "",
        "Ledgerspåret bevarar Process-ID, run metadata och historiska rapportlänkar även när rapporten inte längre ligger i den aktiva hubbvyn.",
        "",
        "| Datum | Roll | Arbetsgivare | Process-ID | Status | Hy-Filter | Rapport | Word | PDF |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for x in sorted(ledger.get("runs", []), key=lambda y: y.get("created_at") or "", reverse=True):
        p = x.get("paths") or {}
        hy = (x.get("hy_filter") or {}).get("status") or "-"
        hlines.append(f"| {str(x.get('created_at') or '')[:10]} | {x.get('title','')} | {x.get('company','')} | {x.get('process_id','')} | {x.get('status','')} | {hy} | [Rapport]({p.get('markdown','#')}) | [Word]({p.get('docx','#')}) | [PDF]({p.get('pdf','#')}) |")
    hist_path.write_text("\n".join(hlines) + "\n", encoding="utf-8")
    return active_path, hist_path


def inject_hrdm_panel(control_room: Path) -> None:
    if not control_room.exists():
        return
    start = "<!-- CAREERHUB:HRDM:START -->"
    end = "<!-- CAREERHUB:HRDM:END -->"
    block = (
        f"{start}\n"
        "## HRDM-analyser\n\n"
        "**[Läs aktiva HRDM-rapporter ->](HRDM_REPORTS.md)**  \n"
        "**[Öppna HRDM-ledgern ->](HRDM_LEDGER.md)**\n\n"
        "Aktiva rapporter försvinner automatiskt från den aktuella rapportvyn när jobbets sista ansökningsdag har passerat. Historiken ligger kvar i ledgern.\n"
        f"{end}"
    )
    text = control_room.read_text(encoding="utf-8")
    if start in text and end in text:
        before = text.split(start, 1)[0].rstrip()
        after = text.split(end, 1)[1].lstrip()
        text = before + "\n\n" + block + "\n\n" + after
    else:
        marker = "## 3 · Sök"
        if marker in text:
            text = text.replace(marker, block + "\n\n" + marker, 1)
        else:
            text = text.rstrip() + "\n\n" + block + "\n"
    control_room.write_text(text, encoding="utf-8")


def maintain_hrdm_reports(root: Path, ledger_path: Path, control_room: Path | None = None) -> int:
    archived = archive_expired_reports(root, ledger_path)
    render_report_indexes(root, ledger_path)
    if control_room:
        inject_hrdm_panel(control_room)
    return archived
